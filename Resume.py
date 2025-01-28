import streamlit as st
import google.generativeai as genai
from pdfminer.high_level import extract_text
from docx import Document
import os
import json
import pandas as pd
from dotenv import load_dotenv
import spacy
from textblob import TextBlob
import asyncio
import nest_asyncio
from openpyxl import load_workbook
import re

nest_asyncio.apply()

nlp = spacy.load("en_core_web_sm")

class ResumeParser:
    def __init__(self):
        load_dotenv()
        api_key = os.getenv("GEMINI_API_KEY")

        if not api_key:
            raise ValueError("GEMINI_API_KEY is not set in the .env file")

        genai.configure(api_key=api_key)
        self.model = genai.GenerativeModel('gemini-pro')

    def parse_file(self, file_path):
        file_extension = os.path.splitext(file_path)[1].lower()

        try:
            if file_extension == '.pdf':
                text = self.parse_pdf(file_path)
            elif file_extension == '.docx':
                text = self.parse_docx(file_path)
            else:
                raise ValueError("Unsupported file format")

            return text
        except Exception as e:
            raise Exception(f"Error parsing file: {str(e)}")

    def parse_pdf(self, file_path):
        try:
            text = extract_text(file_path)
            return text
        except Exception as e:
            raise Exception(f"Error parsing PDF: {str(e)}")

    def parse_docx(self, file_path):
        try:
            doc = Document(file_path)
            text = '\n'.join([paragraph.text for paragraph in doc.paragraphs])
            return text
        except Exception as e:
            raise Exception(f"Error parsing DOCX: {str(e)}")

    def analyze_resume(self, file_path):
        try:
            resume_text = self.parse_file(file_path)

            prompt = f"""
            Extract the following mandatory columns from this resume in JSON format:
            - Name
            - Contact details (as in the resume)
            - University
            - Year of Study
            - Course
            - Discipline
            - CGPA/Percentage
            - Key Skills
            - Gen AI Experience Score
            - AI/ML Experience Score
            - Supporting Information (certifications, internships, projects)

            Resume text:
            {resume_text}
            """

            response = self.model.generate_content(prompt)
            response_text = response.text.strip()

            try:
                extracted_data = json.loads(response_text)
            except json.JSONDecodeError:
                start_idx = response_text.find('{')
                end_idx = response_text.rfind('}') + 1
                if (start_idx != -1 and end_idx != 0):
                    json_str = response_text[start_idx:end_idx]
                    extracted_data = json.loads(json_str)
                else:
                    raise Exception("Could not parse extraction response")

            sentiment_score = self.perform_sentiment_analysis(resume_text)

            extracted_data["Sentiment Score"] = sentiment_score

            experience_scores = self.score_experience(extracted_data.get("Key Skills", ""))
            extracted_data["Experience Scores"] = experience_scores

            if "NER" in extracted_data:
                del extracted_data["NER"]

            if isinstance(extracted_data.get("Contact details"), list):
                extracted_data["Contact details"] = list(set(extracted_data["Contact details"]))

            return extracted_data

        except Exception as e:
            return {"error": str(e)}

    def perform_ner(self, text):
        doc = nlp(text)
        entities = {"names": [], "contact_details": [], "urls": []}

        for ent in doc.ents:
            if ent.label_ == "PERSON":
                entities["names"].append(ent.text)
            elif ent.label_ == "PHONE" or ent.label_ == "EMAIL":
                entities["contact_details"].append(ent.text)

        urls = re.findall(r'https?://(?:www\.)?[\w\-\.]+(?:\.[a-z]{2,})+[/\w\-\.\?=&%]*', text)
        entities["urls"].extend(urls)

        return entities

    def perform_sentiment_analysis(self, text):
        blob = TextBlob(text)
        sentiment = blob.sentiment.polarity
        return sentiment

    def score_experience(self, key_skills):
        ai_ml_keywords = ["AI", "Machine Learning", "Deep Learning", "Neural Networks"]
        gen_ai_keywords = ["Generative AI", "Transformer Models", "GPT", "BERT"]

        ai_ml_score = sum([1 for word in ai_ml_keywords if word in key_skills])
        gen_ai_score = sum([1 for word in gen_ai_keywords if word in key_skills])

        return {"AI/ML Score": ai_ml_score, "Gen AI Score": gen_ai_score}

    def score_role_fit(self, key_skills, required_skills):
        if isinstance(key_skills, list):
            key_skills = " ".join(key_skills)

        match_count = sum([1 for skill in required_skills if skill.lower() in key_skills.lower()])
        score = (match_count / len(required_skills)) * 100
        return score

    def suggest_role(self, role_description, key_skills):
        if not role_description:
            return "Role description is missing. Please provide a description."

        if isinstance(key_skills, list):
            key_skills = " ".join(key_skills)

        match_count = sum([1 for skill in role_description.split() if skill.lower() in key_skills.lower()])

        if match_count > 3:
            return "This resume is a good fit for the specified role."
        elif match_count > 1:
            return "This resume has some relevant experience for the role."
        else:
            return "This resume may not be well-suited for the role based on the key skills."


    def write_to_excel(self, data_list, output_file="extracted_resume_data.xlsx"):
        df = pd.DataFrame(data_list)

        if 'Supporting Information' not in df.columns:
            df['Supporting Information'] = None

        df.to_excel(output_file, index=False, engine='openpyxl')

        wb = load_workbook(output_file)
        ws = wb.active

        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = adjusted_width

        wb.save(output_file)

    async def process_resume_async(self, file_path, required_skills):
        extracted_data = self.analyze_resume(file_path)

        if "Experience Scores" in extracted_data:
            extracted_data["AI/ML Experience Score"] = extracted_data["Experience Scores"]["AI/ML Score"]
            extracted_data["Gen AI Experience Score"] = extracted_data["Experience Scores"]["Gen AI Score"]

        role_fit_score = self.score_role_fit(extracted_data.get("Key Skills", ""), required_skills)
        extracted_data["Role Fit Score"] = role_fit_score

        if 'Supporting Information' not in extracted_data:
            extracted_data['Supporting Information'] = None

        return extracted_data

st.set_page_config(page_title="Generative AI-Powered Resume Analyzer", layout="wide")
st.title("Generative AI-Powered Resume Analyzer")

st.write("Upload one or more resumes to analyze and extract structured information.")
role_description = st.text_area("Enter Role Description (optional)")
required_skills_input = st.text_area("Enter Required Key Skills (comma-separated)", "AI, Machine Learning, Python, SQL")

uploaded_files = st.file_uploader("Upload Resumes", type=None, accept_multiple_files=True)

if uploaded_files:
    if st.button("Start Processing"):
        st.write("Processing resumes...")
        parser = ResumeParser()
        results = []

        required_skills = [skill.strip() for skill in required_skills_input.split(",")]

        for file in uploaded_files:
            with open(file.name, "wb") as f:
                f.write(file.getbuffer())

            extracted_data = parser.analyze_resume(file.name)

            role_fit_score = parser.score_role_fit(extracted_data.get("Key Skills", ""), required_skills_input)
            extracted_data["Role Fit Score"] = role_fit_score

            role_suggestion = parser.suggest_role(role_description, extracted_data.get("Key Skills", ""))
            extracted_data["Suggested Role"] = role_suggestion

            results.append(extracted_data)
            os.remove(file.name)

        output_file = "extracted_resume_data.xlsx"
        parser.write_to_excel(results, output_file)

        st.success("Processing complete!")
        st.download_button(label="Download Excel File", data=open(output_file, "rb").read(), file_name=output_file, mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        for result in results:
            st.json(result)
